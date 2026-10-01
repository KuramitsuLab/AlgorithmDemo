#!/usr/bin/env python3
"""水くみ（バケツ）問題の教材スライド (jug.pptx) を生成する。
docs/jug.html と対になる講義用スライド。
状態 (a,b) をノード・6操作をエッジにした状態空間グラフを描き、
実際に DFS / BFS を走らせた探索結果（訪問・経路）を作図する。
"""
from collections import deque
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ---- カラーパレット (docs/jug.html と統一) ----
BG     = RGBColor(0xF4, 0xF6, 0xFB)
PANEL  = RGBColor(0xFF, 0xFF, 0xFF)
PANEL2 = RGBColor(0xEE, 0xF1, 0xF7)
LINE   = RGBColor(0xD6, 0xDC, 0xE8)
TEXT   = RGBColor(0x1A, 0x20, 0x30)
MUTED  = RGBColor(0x5D, 0x6B, 0x82)
ACCENT = RGBColor(0x2F, 0x7F, 0xD9)
INK    = RGBColor(0x06, 0x12, 0x2A)   # 明色面に乗せる濃文字
ONDARK = RGBColor(0xE6, 0xEC, 0xF5)   # 濃色ノード(訪問/未確定)に乗せる淡文字
NODE   = RGBColor(0x38, 0x44, 0x5F)
FRONT  = RGBColor(0xFF, 0xD2, 0x4A)
VISITED= RGBColor(0x2F, 0x5D, 0x8A)
CURRENT= RGBColor(0xFF, 0x6B, 0x6B)
PATH   = RGBColor(0x4A, 0xDE, 0x80)
GOAL   = RGBColor(0xC0, 0x84, 0xFC)
EDGE   = RGBColor(0x4A, 0x56, 0x72)

FONT = "Meiryo"
MONO = "Consolas"
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

# ====================================================================
# 問題設定と探索
# ====================================================================
A, B, T = 5, 3, 4
def nbrs(s):
    a, b = s
    r = [(A, b), (a, B), (0, b), (a, 0)]
    ab = min(a, B-b); r.append((a-ab, b+ab))
    ba = min(b, A-a); r.append((a+ba, b-ba))
    out, seen = [], {s}
    for x in r:
        if x not in seen:
            seen.add(x); out.append(x)
    return out

def is_goal(s): return s[0] == T or s[1] == T

ALL_STATES = [(a, b) for a in range(A+1) for b in range(B+1)]

# 無向エッジ集合（全状態グラフ）
ALL_EDGES = set()
for s in ALL_STATES:
    for t in nbrs(s):
        ALL_EDGES.add(tuple(sorted((s, t))))
ALL_EDGES = sorted(ALL_EDGES)

def search(bfs):
    start = (0, 0)
    parent = {start: None}
    cont = deque([start]); instr = {start}
    visited = []; visited_set = set()
    while cont:
        n = cont.popleft() if bfs else cont.pop()
        instr.discard(n)
        if n in visited_set: continue
        visited.append(n); visited_set.add(n)
        if is_goal(n):
            path = []; p = n
            while p is not None: path.append(p); p = parent[p]
            return visited_set, path[::-1]
        for x in nbrs(n):
            if x not in visited_set and x not in instr:
                parent[x] = n; instr.add(x); cont.append(x)
    return visited_set, []

bfs_v, bfs_p = search(True)
dfs_v, dfs_p = search(False)

# ====================================================================
# 描画ヘルパ
# ====================================================================
prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

def slide_bg(slide, color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color

def add_text(slide, x, y, w, h, text, size=18, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for j, seg in enumerate(ln.split("**")):
            if seg == "": continue
            r = p.add_run(); r.text = seg
            r.font.size = Pt(size); r.font.name = font
            r.font.color.rgb = ACCENT if (j % 2 == 1) else color
            r.font.bold = bold or (j % 2 == 1)
    return tb

def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    t = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(t, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp

def add_oval(slide, x, y, w, h, fill, ring=None, ring_w=2.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if ring is None: shp.line.fill.background()
    else: shp.line.color.rgb = ring; shp.line.width = Pt(ring_w)
    shp.shadow.inherit = False
    return shp

def add_line(slide, x1, y1, x2, y2, color, width=0.75):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color; cn.line.width = Pt(width)
    cn.shadow.inherit = False
    return cn

def card(slide, x, y, w, h):
    return add_rect(slide, x, y, w, h, PANEL, line=LINE, radius=True)

def section_header(slide, title, subtitle=None, num=None):
    add_rect(slide, 0, 0, EMU_W, Inches(1.15), PANEL)
    add_rect(slide, 0, Inches(1.15), EMU_W, Pt(3), ACCENT)
    if num is not None:
        b = add_rect(slide, Inches(0.45), Inches(0.3), Inches(0.55), Inches(0.55), ACCENT, radius=True)
        p = b.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(num)
        r.font.size = Pt(22); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = INK
    tx = Inches(1.2) if num is not None else Inches(0.5)
    add_text(slide, tx, Inches(0.18), Inches(11), Inches(0.6), title, size=28, bold=True)
    if subtitle:
        add_text(slide, tx, Inches(0.72), Inches(11.5), Inches(0.4), subtitle, size=14, color=MUTED)

def draw_graph(slide, ox, oy, gw, gh, visited=None, path=None,
               show_edges=True, show_undiscovered=True, label=True, rscale=0.30):
    """状態空間グラフを描く。visited/path は (a,b) の集合/リスト。"""
    visited = visited or set()
    pathset = set(path or [])
    path_edges = set()
    if path:
        for i in range(len(path)-1):
            path_edges.add(tuple(sorted((path[i], path[i+1]))))
    colw = gw / A
    rowh = gh / B
    rad = int(min(colw, rowh) * rscale)
    def cx(a): return int(ox + a*colw)
    def cy(b): return int(oy + (B-b)*rowh)
    # 1) 通常エッジ（薄く）
    if show_edges:
        for (s, t) in ALL_EDGES:
            if (s, t) in path_edges: continue
            add_line(slide, cx(s[0]), cy(s[1]), cx(t[0]), cy(t[1]), EDGE, 0.75)
    # 2) 経路エッジ（緑・太）
    for (s, t) in path_edges:
        add_line(slide, cx(s[0]), cy(s[1]), cx(t[0]), cy(t[1]), PATH, 2.75)
    # 3) ノード
    for (a, b) in ALL_STATES:
        x, y = cx(a)-rad, cy(b)-rad
        in_path = (a, b) in pathset
        in_vis = (a, b) in visited
        if not in_path and not in_vis and not show_undiscovered:
            continue
        if in_path:    fill = PATH
        elif in_vis:   fill = VISITED
        else:          fill = NODE
        ring = None; rw = 2.5
        if (a, b) == (0, 0):     ring = RGBColor(0x28,0x37,0x50); rw = 2.5
        elif is_goal((a, b)):    ring = GOAL; rw = 2.5
        elif not in_path and not in_vis:
            ring = RGBColor(0x4A,0x56,0x72); rw = 1.0   # 未発見は淡い枠
        ov = add_oval(slide, Emu(x), Emu(y), Emu(2*rad), Emu(2*rad), fill, ring=ring, ring_w=rw)
        if label and rad > 12:
            tf = ov.text_frame; tf.word_wrap = False
            tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            rn = p.add_run(); rn.text = f"{a},{b}"
            rn.font.size = Pt(max(7, min(13, round(rad/12700*0.62))))   # rad(EMU)→pt
            rn.font.bold = True; rn.font.name = FONT
            rn.font.color.rgb = INK if fill in (PATH,) else ONDARK
    return rad

def legend(slide, x, y, items):
    cx = x
    for label_t, col in items:
        add_oval(slide, cx, y + Inches(0.02), Inches(0.17), Inches(0.17), col)
        add_text(slide, cx + Inches(0.23), y - Inches(0.03), Inches(2.2), Inches(0.3),
                 label_t, size=11, color=MUTED)
        cx += Inches(0.23) + Inches(len(label_t) * 0.16 + 0.4)

# ====================================================================
# 1. タイトル
# ====================================================================
s = prs.slides.add_slide(BLANK); slide_bg(s)
draw_graph(s, Inches(8.6), Inches(2.3), Inches(3.7), Inches(2.3),
           visited=bfs_v, path=bfs_p, rscale=0.28)
add_text(s, Inches(0.7), Inches(1.9), Inches(8), Inches(1.0),
         "水くみ問題で学ぶ探索", size=46, bold=True)
add_text(s, Inches(0.7), Inches(3.0), Inches(8), Inches(0.6),
         "状態グラフを DFS と BFS でたどる", size=24, color=ACCENT)
add_text(s, Inches(0.7), Inches(3.8), Inches(8), Inches(0.5),
         f"5L と 3L のバケツで、ちょうど {T}L を作れる？", size=16, color=MUTED)
add_text(s, Inches(0.7), Inches(6.7), Inches(8), Inches(0.4),
         "対応ビジュアライザ:  docs/jug.html", size=13, color=MUTED, font=MONO)

# ====================================================================
# 2. 問題と「状態・操作」
# ====================================================================
s = prs.slides.add_slide(BLANK); slide_bg(s)
section_header(s, "バケツ問題とは", f"5L と 3L のバケツだけで、ちょうど {T}L を量る", num=1)
card(s, Inches(0.5), Inches(1.5), Inches(6.3), Inches(5.2))
add_text(s, Inches(0.75), Inches(1.65), Inches(5.8), Inches(0.4), "ルール", size=18, bold=True, color=ACCENT)
add_text(s, Inches(0.75), Inches(2.15), Inches(5.8), Inches(1.2),
         "・ 目盛りは無い。できるのは下の操作だけ\n"
         "・ バケツ A=5L, バケツ B=3L\n"
         f"・ スタートは (0,0)。どちらかが {T}L になれば成功",
         size=15, line_spacing=1.3)
add_text(s, Inches(0.75), Inches(3.5), Inches(5.8), Inches(0.4), "状態の表し方", size=18, bold=True, color=ACCENT)
add_text(s, Inches(0.75), Inches(4.0), Inches(5.8), Inches(0.7),
         "今の水量を **(Aの水, Bの水)** と書く。\n"
         f"例: (0,0) 空っぽ → 目標 (例 {T},0) を目指す",
         size=15, line_spacing=1.3)
add_text(s, Inches(0.75), Inches(4.9), Inches(5.8), Inches(0.4), "6つの操作", size=18, bold=True, color=ACCENT)
add_text(s, Inches(0.75), Inches(5.35), Inches(5.8), Inches(1.2),
         "① A を満たす ② B を満たす\n"
         "③ A を空にする ④ B を空にする\n"
         "⑤ A→B に移す ⑥ B→A に移す",
         size=15, line_spacing=1.35)
# 右: 状態グラフ全体
card(s, Inches(7.1), Inches(1.5), Inches(5.7), Inches(5.2))
add_text(s, Inches(7.35), Inches(1.62), Inches(5.2), Inches(0.4),
         "すべての状態とつながり（状態空間グラフ）", size=13, color=MUTED)
draw_graph(s, Inches(8.0), Inches(2.5), Inches(3.9), Inches(2.9), show_edges=True)
add_text(s, Inches(7.35), Inches(5.7), Inches(5.2), Inches(0.9),
         f"丸 = 状態（全 {len(ALL_STATES)} 個）、線 = 操作でのつながり。\n"
         "濃い枠 = スタート (0,0)、紫枠 = 目標の状態。", size=12, color=MUTED, line_spacing=1.3)

# ====================================================================
# 3. 探索の枠組み（DFS と BFS）
# ====================================================================
s = prs.slides.add_slide(BLANK); slide_bg(s)
section_header(s, "迷路と同じ「グラフ探索」", "状態をたどってスタートからゴールへ。違いは“次にどれを選ぶか”", num=2)
cols = [
    ("DFS", "スタック", "最後に入れた候補を選ぶ", "1本の手順を奥まで試す", CURRENT),
    ("BFS", "キュー",   "最初に入れた候補を選ぶ", "手数の少ない順に広げる", ACCENT),
]
for i, (name, ds, how, char, col) in enumerate(cols):
    x = Inches(0.5) + i*Inches(6.4)
    card(s, x, Inches(1.6), Inches(6.0), Inches(2.7))
    add_rect(s, x, Inches(1.6), Inches(6.0), Inches(0.65), PANEL2, radius=True)
    add_text(s, x, Inches(1.68), Inches(6.0), Inches(0.5), name, size=22, bold=True, color=col, align=PP_ALIGN.CENTER)
    add_text(s, x+Inches(0.3), Inches(2.5), Inches(5.4), Inches(0.4), "データ構造", size=12, color=MUTED)
    add_text(s, x+Inches(0.3), Inches(2.85), Inches(5.4), Inches(0.4), ds, size=16, bold=True)
    add_text(s, x+Inches(0.3), Inches(3.3), Inches(5.4), Inches(0.4), "選び方 / 性格", size=12, color=MUTED)
    add_text(s, x+Inches(0.3), Inches(3.65), Inches(5.4), Inches(0.6), how + "\n→ " + char, size=15, line_spacing=1.2)
add_rect(s, Inches(0.5), Inches(4.6), Inches(12.3), Inches(2.0), PANEL2, line=ACCENT, radius=True)
add_text(s, Inches(0.8), Inches(4.75), Inches(11.7), Inches(0.5),
         "共通の手順", size=16, bold=True, color=ACCENT)
add_text(s, Inches(0.8), Inches(5.25), Inches(11.7), Inches(1.2),
         "1. スタート (0,0) を入れる   2. 1つ取り出して訪問   "
         "3. そこから操作で行ける状態（隣）を入れる   4. ゴールに着くまで繰り返す\n"
         "取り出し方が **スタック（DFS）か キュー（BFS）か** だけで、探索の順番と得られる手数が変わる。",
         size=15, line_spacing=1.4)

# ====================================================================
# アルゴリズム結果スライド
# ====================================================================
def algo_slide(num, name_jp, name_en, idea, visited, path, struct, order, shortest, note, col):
    s = prs.slides.add_slide(BLANK); slide_bg(s)
    section_header(s, name_jp, name_en, num=num)
    # 左: 解説
    card(s, Inches(0.5), Inches(1.5), Inches(5.5), Inches(5.4))
    add_text(s, Inches(0.75), Inches(1.65), Inches(5.0), Inches(0.4), "考え方", size=17, bold=True, color=ACCENT)
    add_text(s, Inches(0.75), Inches(2.1), Inches(5.0), Inches(0.9), idea, size=14, line_spacing=1.25)
    add_text(s, Inches(0.75), Inches(3.0), Inches(5.0), Inches(0.4), "見つけた手順", size=17, bold=True, color=ACCENT)
    steps_txt = "\n".join(f"{i}. ({a},{b})" for i, (a, b) in enumerate(path))
    add_text(s, Inches(0.75), Inches(3.45), Inches(5.0), Inches(2.6), steps_txt, size=13, font=MONO, line_spacing=1.2, color=PATH)
    add_text(s, Inches(0.75), Inches(6.2), Inches(5.0), Inches(0.7),
             "ポイント: " + note, size=12.5, color=FRONT, line_spacing=1.2)
    # 右上: グラフ
    card(s, Inches(6.2), Inches(1.5), Inches(6.6), Inches(3.7))
    add_text(s, Inches(6.45), Inches(1.6), Inches(6.1), Inches(0.35),
             f"探索結果 — 訪問 {len(visited)} 状態 / 手数 {len(path)-1}", size=12, color=MUTED)
    draw_graph(s, Inches(7.4), Inches(2.35), Inches(4.2), Inches(2.3),
               visited=visited, path=path, rscale=0.30)
    legend(s, Inches(6.45), Inches(4.85),
           [("訪問", VISITED), ("経路", PATH), ("目標", GOAL)])
    # 右下: 性質
    card(s, Inches(6.2), Inches(5.4), Inches(6.6), Inches(1.5))
    info = [("データ構造", struct), ("探索順", order), ("最少手数", shortest)]
    cy = Inches(5.55)
    for k, v in info:
        add_text(s, Inches(6.45), cy, Inches(2.2), Inches(0.3), k, size=12, color=MUTED)
        add_text(s, Inches(8.5), cy, Inches(4.1), Inches(0.3), v, size=12, color=TEXT, font=MONO)
        cy += Inches(0.32)

algo_slide(
    3, "DFS（深さ優先探索）", "Depth-First Search",
    "1つの操作を選んだら、行き詰まるまで\nどんどん奥の状態へ進む。",
    dfs_v, dfs_p, "スタック", "奥へ一直線", "× 保証なし",
    f"訪問は少ない（{len(dfs_v)}）が、手数は {len(dfs_p)-1} 手と遠回りになった。", CURRENT)

algo_slide(
    4, "BFS（幅優先探索）", "Breadth-First Search",
    "スタートから1手で行ける状態、2手で\n行ける状態…と近い順に広げる。",
    bfs_v, bfs_p, "キュー", "近い手数から順", "○ 最少手数を保証",
    f"訪問は多め（{len(bfs_v)}）だが、必ず最少の {len(bfs_p)-1} 手で到達できる。", ACCENT)

# ====================================================================
# 5. 比較
# ====================================================================
s = prs.slides.add_slide(BLANK); slide_bg(s)
section_header(s, "DFS と BFS を比べる", "同じ問題でも、得られる手数が違う", num=5)
rows = [
    ("", "DFS", "BFS"),
    ("データ構造", "スタック", "キュー"),
    ("探索順", "奥へ一直線", "近い手数から"),
    ("訪問した状態", str(len(dfs_v)), str(len(bfs_v))),
    ("見つけた手数", f"{len(dfs_p)-1} 手", f"{len(bfs_p)-1} 手"),
    ("最少手数", "× 保証なし", "○ 保証あり"),
]
nrows, ncols = len(rows), 3
tbl = s.shapes.add_table(nrows, ncols, Inches(0.5), Inches(1.55), Inches(6.2), Inches(4.0)).table
tbl.columns[0].width = Inches(2.6); tbl.columns[1].width = Inches(1.8); tbl.columns[2].width = Inches(1.8)
for r in range(nrows):
    tbl.rows[r].height = Inches(0.66)
    for c in range(ncols):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.1); cell.margin_right = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0 or c == 0:
            cell.fill.fore_color.rgb = PANEL2
        else:
            cell.fill.fore_color.rgb = PANEL if r % 2 else PANEL2
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        rn = p.add_run(); rn.text = rows[r][c]
        rn.font.size = Pt(14); rn.font.name = FONT
        rn.font.bold = (r == 0 or c == 0)
        if r == 0 and c == 1: rn.font.color.rgb = CURRENT
        elif r == 0 and c == 2: rn.font.color.rgb = ACCENT
        elif c == 0: rn.font.color.rgb = MUTED
        elif "○" in rows[r][c]: rn.font.color.rgb = PATH
        elif "×" in rows[r][c]: rn.font.color.rgb = MUTED
        else: rn.font.color.rgb = TEXT
# 右: 2グラフ
card(s, Inches(7.0), Inches(1.55), Inches(5.8), Inches(4.0))
add_text(s, Inches(7.2), Inches(1.65), Inches(2.6), Inches(0.35), "DFS（8手・遠回り）", size=12, color=CURRENT)
draw_graph(s, Inches(7.6), Inches(2.2), Inches(2.0), Inches(1.3), visited=dfs_v, path=dfs_p, label=False, rscale=0.26)
add_text(s, Inches(10.1), Inches(1.65), Inches(2.6), Inches(0.35), "BFS（6手・最少）", size=12, color=ACCENT)
draw_graph(s, Inches(10.5), Inches(2.2), Inches(2.0), Inches(1.3), visited=bfs_v, path=bfs_p, label=False, rscale=0.26)
add_text(s, Inches(7.2), Inches(3.9), Inches(5.4), Inches(1.5),
         "BFS は手数の少ない状態から調べるので、**必ず最少手数**。\n"
         "DFS は奥へ突き進むため、訪問は少なくても **遠回りの手順**になりやすい。",
         size=14, color=TEXT, line_spacing=1.3)
add_text(s, Inches(0.5), Inches(5.9), Inches(6.2), Inches(0.9),
         f"※ どちらも「Aが{T}L」を達成。BFS は (4,3)、DFS は (4,0) に到達した。",
         size=12.5, color=MUTED, line_spacing=1.2)

# ====================================================================
# 6. まとめ & 演習
# ====================================================================
s = prs.slides.add_slide(BLANK); slide_bg(s)
section_header(s, "まとめと演習", "手を動かして確かめよう", num=6)
card(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(3.1))
add_text(s, Inches(0.75), Inches(1.65), Inches(5.5), Inches(0.4), "今日のポイント", size=18, bold=True, color=ACCENT)
add_text(s, Inches(0.75), Inches(2.2), Inches(5.5), Inches(2.3),
         "・ パズルも **状態（ノード）と操作（エッジ）** のグラフ\n"
         "・ 解くこと＝ゴール状態への **経路探索**\n"
         "・ **BFS** は最少手数を保証、**DFS** はしない\n"
         "・ **DFS** は訪問が少ないが遠回りもある\n"
         "・ 迷路でも水くみでも、同じ枠組みで解ける",
         size=15, line_spacing=1.35)
card(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(3.1))
add_text(s, Inches(7.05), Inches(1.65), Inches(5.5), Inches(0.4), "演習（ビジュアライザで）", size=18, bold=True, color=ACCENT)
add_text(s, Inches(7.05), Inches(2.2), Inches(5.5), Inches(2.3),
         "1. DFS と BFS の「手数」を見比べよう\n"
         "2. バケツを A=7, B=5 に変えると\n   状態（ノード）は何個に増える？\n"
         "3. 目標を変えると到達できる？\n   できない目標もある（例: 偶数だけ 等）",
         size=15, line_spacing=1.3)
add_rect(s, Inches(0.5), Inches(4.9), Inches(12.3), Inches(1.5), PANEL2, line=ACCENT, radius=True)
add_text(s, Inches(0.8), Inches(5.05), Inches(11.7), Inches(0.5),
         "▶ 付属ビジュアライザで体験", size=18, bold=True, color=ACCENT)
add_text(s, Inches(0.8), Inches(5.55), Inches(11.7), Inches(0.8),
         "docs/jug.html をブラウザで開く → A・B・目標を設定 → 「次へ」で1手ずつ状態が増えていく / "
         "◀戻るで逆再生 / DFS↔BFS を切り替えて手数を比較",
         size=14, color=TEXT, font=MONO, line_spacing=1.2)
add_text(s, Inches(0.5), Inches(6.95), Inches(12.3), Inches(0.4),
         "おしまい — 状態グラフが育つ様子を自分の手で動かそう！", size=14, color=MUTED, align=PP_ALIGN.CENTER)

out = "/Users/kimio/AI/Algo/jug.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
print(f"BFS: 訪問{len(bfs_v)} 手数{len(bfs_p)-1}  経路 {bfs_p}")
print(f"DFS: 訪問{len(dfs_v)} 手数{len(dfs_p)-1}  経路 {dfs_p}")
print(f"全状態数 {len(ALL_STATES)}  全エッジ {len(ALL_EDGES)}")
