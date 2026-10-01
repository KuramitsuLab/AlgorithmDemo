#!/usr/bin/env python3
"""ビームサーチの教材スライド (beam.pptx) を生成する。
docs/beam.html と対になる講義用スライド。
おもちゃの言語モデルで実際に探索木を構築し、貪欲(k=1)とビーム(k=2)を比較する。
"""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

BG=RGBColor(0xF4,0xF6,0xFB); PANEL=RGBColor(0xFF,0xFF,0xFF); PANEL2=RGBColor(0xEE,0xF1,0xF7)
LINE=RGBColor(0xD6,0xDC,0xE8); TEXT=RGBColor(0x1A,0x20,0x30); MUTED=RGBColor(0x5D,0x6B,0x82)
ACCENT=RGBColor(0x2F,0x7F,0xD9); CAND=RGBColor(0xFF,0xD2,0x4A); BEAM=RGBColor(0x2F,0x7F,0xD9)
INK=RGBColor(0x06,0x12,0x2A)        # 明色面(候補/ビーム/最良ノード)に乗せる濃文字
ONDARK=RGBColor(0xE6,0xEC,0xF5)     # 濃色(枝刈り)ノードに乗せる淡文字
PRUNED=RGBColor(0x39,0x43,0x5A); BEST=RGBColor(0x4A,0xDE,0x80); ENDC=RGBColor(0xC0,0x84,0xFC)
FONT="Meiryo"; MONO="Consolas"
EMU_W=Inches(13.333); EMU_H=Inches(7.5)

# ---- おもちゃの言語モデル（trap：貪欲が失敗する例）----
TABLE={
 "":{"猫":0.55,"犬":0.45},"猫":{"は":0.9,"が":0.1},"犬":{"は":0.9,"が":0.1},
 "猫 は":{"寝る":0.5,"鳴く":0.5},"犬 は":{"走る":0.95,"吠える":0.05},
 "猫 は 寝る":{"。":1.0},"猫 は 鳴く":{"。":1.0},"犬 は 走る":{"。":1.0},"犬 は 吠える":{"。":1.0},
 "猫 が":{"好き":1.0},"犬 が":{"好き":1.0},"猫 が 好き":{"。":1.0},"犬 が 好き":{"。":1.0},
}

def build_tree(K, maxlen=4):
    nodes=[]; nid=0
    nodes.append({"id":0,"parent":-1,"token":"▶","seq":[],"prob":1.0,"lp":0.0,"depth":0,"state":"beam"}); nid=1
    beam=[0]
    for depth in range(1,maxlen+2):
        cands=[]; allend=True
        for bid in beam:
            node=nodes[bid]
            if node["seq"] and node["token"]=="。":
                cands.append(bid); continue
            allend=False
            dist=TABLE.get(" ".join(node["seq"]),{"。":1.0})
            for tok,p in dist.items():
                nodes.append({"id":nid,"parent":bid,"token":tok,"seq":node["seq"]+[tok],
                    "prob":node["prob"]*p,"lp":node["lp"]+math.log(p),"depth":depth,"state":"cand"})
                cands.append(nid); nid+=1
        if allend: break
        order=sorted(cands,key=lambda i:nodes[i]["lp"],reverse=True)
        keep=set(order[:K])
        for cid in cands:
            if nodes[cid]["state"]=="cand":
                nodes[cid]["state"]="beam" if cid in keep else "pruned"
        beam=order[:K]
        if all(nodes[b]["token"]=="。" for b in beam): break
    best=sorted(beam,key=lambda i:nodes[i]["lp"],reverse=True)[0]
    p=best
    while p!=-1:
        if nodes[p]["state"]!="pruned": nodes[p]["state"]="best"
        p=nodes[p]["parent"]
    return nodes,best

prs=Presentation(); prs.slide_width=EMU_W; prs.slide_height=EMU_H
BLANK=prs.slide_layouts[6]

def slide_bg(s,c=BG): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def add_text(s,x,y,w,h,t,size=18,color=TEXT,bold=False,align=PP_ALIGN.LEFT,font=FONT,
             anchor=MSO_ANCHOR.TOP,ls=1.1):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,ln in enumerate(t.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.line_spacing=ls
        for j,seg in enumerate(ln.split("**")):
            if seg=="":continue
            r=p.add_run(); r.text=seg; r.font.size=Pt(size); r.font.name=font
            r.font.color.rgb=ACCENT if j%2 else color; r.font.bold=bold or bool(j%2)
    return tb
def add_rect(s,x,y,w,h,fill,line=None,radius=False):
    t=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh=s.shapes.add_shape(t,x,y,w,h); sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(1.0)
    sh.shadow.inherit=False; return sh
def add_line(s,x1,y1,x2,y2,color,width=1.0):
    cn=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,int(x1),int(y1),int(x2),int(y2))
    cn.line.color.rgb=color; cn.line.width=Pt(width); cn.shadow.inherit=False; return cn
def card(s,x,y,w,h): return add_rect(s,x,y,w,h,PANEL,line=LINE,radius=True)
def header(s,title,sub=None,num=None):
    add_rect(s,0,0,EMU_W,Inches(1.15),PANEL); add_rect(s,0,Inches(1.15),EMU_W,Pt(3),ACCENT)
    if num is not None:
        b=add_rect(s,Inches(0.45),Inches(0.3),Inches(0.55),Inches(0.55),ACCENT,radius=True)
        p=b.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=str(num); r.font.size=Pt(22); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=INK
    tx=Inches(1.2) if num is not None else Inches(0.5)
    add_text(s,tx,Inches(0.18),Inches(11),Inches(0.6),title,size=28,bold=True)
    if sub: add_text(s,tx,Inches(0.72),Inches(11.5),Inches(0.4),sub,size=14,color=MUTED)

STATE_FILL={"beam":BEAM,"best":BEST,"pruned":PRUNED,"cand":CAND}
def draw_tree(s, ox,oy,w,h, nodes, label=True, nbw=0.92, nbh=0.40):
    ox,oy,w,h=int(ox),int(oy),int(w),int(h)
    maxd=max(n["depth"] for n in nodes)
    bydepth={}
    for n in nodes: bydepth.setdefault(n["depth"],[]).append(n)
    colw=w/max(1,maxd)
    pos={}
    for d in range(maxd+1):
        ids=bydepth.get(d,[]); cnt=len(ids); gap=h/max(1,cnt)
        for i,n in enumerate(ids):
            pos[n["id"]]=(ox + d*colw, int(oy + gap*(i+0.5)))
    nbw_e=Inches(nbw); nbh_e=Inches(nbh)
    # エッジ
    for n in nodes:
        if n["parent"]<0: continue
        a=pos[n["parent"]]; b=pos[n["id"]]
        best=n["state"]=="best"
        col=BEST if best else (RGBColor(0x2A,0x32,0x45) if n["state"]=="pruned" else RGBColor(0x46,0x52,0x6e))
        add_line(s,a[0],a[1],b[0],b[1],col,2.4 if best else 1.0)
    # ノード
    for n in nodes:
        x,y=pos[n["id"]]
        fill=STATE_FILL.get(n["state"],PANEL2)
        box=add_rect(s,Emu(int(x-nbw_e/2)),Emu(int(y-nbh_e/2)),nbw_e,nbh_e,fill,radius=True)
        if n["token"]=="。": box.line.color.rgb=ENDC; box.line.width=Pt(2.0)
        if label:
            tf=box.text_frame; tf.word_wrap=False
            tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
            p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
            r=p.add_run(); r.text=n["token"]
            r.font.size=Pt(12); r.font.bold=True; r.font.name=FONT
            r.font.color.rgb=INK if n["state"] in ("beam","best","cand") else ONDARK
            if n["parent"]>=0:
                add_text(s,Emu(int(x-nbw_e/2)),Emu(int(y+nbh_e/2-Inches(0.02))),nbw_e,Inches(0.26),
                         f"{n['prob']*100:.0f}%",size=9,
                         color=(RGBColor(0x9A,0xA6,0xBD) if n["state"]=="pruned" else MUTED),
                         align=PP_ALIGN.CENTER,font=MONO)
    return pos

# 探索木を構築
nodes2,best2=build_tree(2)
nodes1,best1=build_tree(1)
def seqstr(nodes,bid): return "".join(nodes[bid]["seq"])
def prob(nodes,bid): return nodes[bid]["prob"]

# ====================================================================
# 1. タイトル
# ====================================================================
s=prs.slides.add_slide(BLANK); slide_bg(s)
draw_tree(s, Inches(7.4),Inches(1.7),Inches(5.3),Inches(4.2), nodes2, nbw=0.85)
add_text(s,Inches(0.7),Inches(2.0),Inches(6.4),Inches(1.0),"ビームサーチ",size=48,bold=True)
add_text(s,Inches(0.7),Inches(3.1),Inches(6.4),Inches(0.6),"LLM が文章を生成する仕組み",size=22,color=ACCENT)
add_text(s,Inches(0.7),Inches(3.9),Inches(6.4),Inches(0.5),"確率の高い候補だけ残して、続きを探す",size=16,color=MUTED)
add_text(s,Inches(0.7),Inches(6.7),Inches(6),Inches(0.4),"対応ビジュアライザ:  docs/beam.html",size=13,color=MUTED,font=MONO)

# ====================================================================
# 2. LLM は1語ずつ確率で生成
# ====================================================================
s=prs.slides.add_slide(BLANK); slide_bg(s)
header(s,"LLM は文章を1語ずつ作る","次の単語を「確率」で予想して、つないでいく",num=1)
card(s,Inches(0.5),Inches(1.5),Inches(6.3),Inches(5.2))
add_text(s,Inches(0.75),Inches(1.65),Inches(5.8),Inches(0.4),"次の単語は確率で決まる",size=18,bold=True,color=ACCENT)
add_text(s,Inches(0.75),Inches(2.15),Inches(5.8),Inches(1.3),
         "「犬 は」の次に来る単語を、モデルは\n"
         "確率で予想する：\n"
         "  走る 95% ／ 吠える 5% …",size=15,ls=1.35)
add_text(s,Inches(0.75),Inches(3.5),Inches(5.8),Inches(0.4),"全部試すと爆発する",size=18,bold=True,color=ACCENT)
add_text(s,Inches(0.75),Inches(4.0),Inches(5.8),Inches(1.4),
         "語彙が1万語あれば、3語先まで全部試すと\n"
         "1万 × 1万 × 1万 = 1兆通り。\n"
         "**すべては調べられない** → 賢く絞る必要がある。",size=15,ls=1.35)
add_text(s,Inches(0.75),Inches(5.5),Inches(5.8),Inches(1.1),
         "文の良さ＝**各単語の確率の掛け算（累積確率）**。\n"
         "これが大きい文ほど「もっともらしい」。",size=14,color=MUTED,ls=1.35)
# 右：1ステップの確率分布イメージ
card(s,Inches(7.1),Inches(1.5),Inches(5.7),Inches(5.2))
add_text(s,Inches(7.35),Inches(1.65),Inches(5.2),Inches(0.4),"「犬 は ___」次の単語の確率",size=14,bold=True)
bars=[("走る",0.95,BEAM),("吠える",0.05,PRUNED)]
by=Inches(2.4)
for tok,p,col in bars:
    add_text(s,Inches(7.35),by,Inches(1.3),Inches(0.4),tok,size=15)
    add_rect(s,Inches(8.7),by+Inches(0.03),Emu(int(Inches(3.6)*p)),Inches(0.34),col,radius=True)
    add_text(s,Emu(int(Inches(8.7)+Inches(3.6)*p+Inches(0.05))),by,Inches(1.0),Inches(0.4),
             f"{p*100:.0f}%",size=13,color=MUTED)
    by+=Inches(0.7)
add_text(s,Inches(7.35),Inches(4.0),Inches(5.2),Inches(2.4),
         "毎ステップこの予想を出し、単語を選んで\n"
         "つなげていく。\n\n"
         "「どれを選び、どれを残すか」の戦略が\n"
         "**ビームサーチ**。",size=15,color=TEXT,ls=1.4)

# ====================================================================
# 3. ビームサーチの手順
# ====================================================================
s=prs.slides.add_slide(BLANK); slide_bg(s)
header(s,"ビームサーチの手順","毎ステップ、上位 k 本（ビーム幅）だけ残して進む",num=2)
steps=[("① 展開","今残っている各文から、次の単語の候補をすべて作る",CAND),
       ("② 採点","それぞれの累積確率（確率の掛け算）を計算する",ACCENT),
       ("③ 選択","上位 k 本だけ残し、残りは捨てる（枝刈り）",BEAM),
       ("④ 繰り返し","文末「。」になるまで①〜③を繰り返す",BEST)]
for i,(t,d,col) in enumerate(steps):
    y=Inches(1.55)+i*Inches(0.92)
    add_rect(s,Inches(0.5),y,Inches(6.2),Inches(0.78),PANEL2,radius=True)
    add_text(s,Inches(0.7),y+Inches(0.1),Inches(1.5),Inches(0.6),t,size=17,bold=True,color=col)
    add_text(s,Inches(2.2),y+Inches(0.13),Inches(4.3),Inches(0.6),d,size=13.5,ls=1.15,anchor=MSO_ANCHOR.MIDDLE)
card(s,Inches(7.0),Inches(1.55),Inches(5.8),Inches(4.95))
add_text(s,Inches(7.25),Inches(1.65),Inches(5.3),Inches(0.4),"k=2 で広がる探索木（緑=最良の文）",size=13,color=MUTED)
draw_tree(s, Inches(7.45),Inches(2.2),Inches(5.0),Inches(3.9), nodes2, nbw=0.82)
add_text(s,Inches(0.5),Inches(5.5),Inches(6.2),Inches(1.2),
         "青 = 残したビーム、灰 = 枝刈り、緑 = 最終的に選ばれた文。\n"
         "ビーム幅 k を大きくすると探索は広く・丁寧になるが、計算は重くなる。",size=13.5,color=MUTED,ls=1.35)

# ====================================================================
# 4. 貪欲 vs ビーム（わな）
# ====================================================================
s=prs.slides.add_slide(BLANK); slide_bg(s)
header(s,"貪欲法（k=1）の落とし穴","毎回1位を選ぶと、全体ではもっともらしくない文になることがある",num=3)
# 左：k=1
card(s,Inches(0.5),Inches(1.5),Inches(6.1),Inches(4.4))
add_text(s,Inches(0.7),Inches(1.62),Inches(5.7),Inches(0.4),"貪欲法 k=1（毎回1位だけ）",size=15,bold=True,color=RGBColor(0xFF,0x6B,0x6B))
draw_tree(s, Inches(0.75),Inches(2.15),Inches(5.5),Inches(3.3), nodes1, nbw=0.82)
add_text(s,Inches(0.7),Inches(5.45),Inches(5.7),Inches(0.4),
         f"結果：「{seqstr(nodes1,best1)}」 累積 {prob(nodes1,best1)*100:.1f}%",size=14,bold=True,color=RGBColor(0xFF,0x6B,0x6B))
# 右：k=2
card(s,Inches(6.7),Inches(1.5),Inches(6.1),Inches(4.4))
add_text(s,Inches(6.9),Inches(1.62),Inches(5.7),Inches(0.4),"ビーム k=2（上位2本を保持）",size=15,bold=True,color=ACCENT)
draw_tree(s, Inches(6.95),Inches(2.15),Inches(5.5),Inches(3.3), nodes2, nbw=0.82)
add_text(s,Inches(6.9),Inches(5.45),Inches(5.7),Inches(0.4),
         f"結果：「{seqstr(nodes2,best2)}」 累積 {prob(nodes2,best2)*100:.1f}%",size=14,bold=True,color=BEST)
add_rect(s,Inches(0.5),Inches(6.05),Inches(12.3),Inches(1.0),PANEL2,line=ACCENT,radius=True)
add_text(s,Inches(0.8),Inches(6.18),Inches(11.7),Inches(0.8),
         f"k=1 は最初に高確率の「猫」(55%)につられ、全体では {prob(nodes1,best1)*100:.1f}% の文止まり。"
         f"k=2 は「犬」も残したことで、より良い {prob(nodes2,best2)*100:.1f}% の文を発見。"
         "＝ **少し広く残すと、より良い文に届く**。",size=14,ls=1.35)

# ====================================================================
# 5. まとめ & 演習
# ====================================================================
s=prs.slides.add_slide(BLANK); slide_bg(s)
header(s,"まとめと演習","手を動かして確かめよう",num=4)
card(s,Inches(0.5),Inches(1.5),Inches(6.0),Inches(3.1))
add_text(s,Inches(0.75),Inches(1.65),Inches(5.5),Inches(0.4),"今日のポイント",size=18,bold=True,color=ACCENT)
add_text(s,Inches(0.75),Inches(2.2),Inches(5.5),Inches(2.3),
         "・ LLM は次の単語を**確率**で予想して生成\n"
         "・ 全部試すと**爆発**するので候補を絞る\n"
         "・ 毎ステップ**上位 k 本**だけ残す＝ビーム\n"
         "・ **k=1 は貪欲法**：速いが最良を逃すことも\n"
         "・ k を上げると丁寧だが計算は重くなる",size=15,ls=1.35)
card(s,Inches(6.8),Inches(1.5),Inches(6.0),Inches(3.1))
add_text(s,Inches(7.05),Inches(1.65),Inches(5.5),Inches(0.4),"演習（ビジュアライザで）",size=18,bold=True,color=ACCENT)
add_text(s,Inches(7.05),Inches(2.2),Inches(5.5),Inches(2.3),
         "1. k=1 と k=2 で出来る文と累積確率を\n   見比べよう（わなに注目）\n"
         "2. k を 3,4 と上げると結果は変わる？\n"
         "3. 枝刈りで「捨てた」候補が、実は\n   良かったケースはある？",size=15,ls=1.3)
add_rect(s,Inches(0.5),Inches(4.9),Inches(12.3),Inches(1.5),PANEL2,line=ACCENT,radius=True)
add_text(s,Inches(0.8),Inches(5.05),Inches(11.7),Inches(0.5),"▶ 付属ビジュアライザで体験",size=18,bold=True,color=ACCENT)
add_text(s,Inches(0.8),Inches(5.55),Inches(11.7),Inches(0.8),
         "docs/beam.html を開く → ビーム幅 k を設定 → 「次へ」で1語ずつ展開・枝刈り / "
         "残ったビームと累積確率、最終的な最良文をその場で確認",size=14,color=TEXT,font=MONO,ls=1.2)
add_text(s,Inches(0.5),Inches(6.95),Inches(12.3),Inches(0.4),
         "「どれを残すか」の工夫が、AI の文章生成の質を決めている。",size=14,color=MUTED,align=PP_ALIGN.CENTER)

out="/Users/kimio/AI/Algo/beam.pptx"
prs.save(out)
print("saved:",out,"slides:",len(prs.slides._sldIdLst))
print(f"k=1: {seqstr(nodes1,best1)} {prob(nodes1,best1)*100:.1f}%")
print(f"k=2: {seqstr(nodes2,best2)} {prob(nodes2,best2)*100:.1f}%")
